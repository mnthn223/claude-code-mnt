"""変額終身保険（プルデンシャル生命・終身型）の払込満了後の減額による推移の試算表を作る。

使い方: uv run --no-project --with openpyxl python build_xlsx.py qx_m_death.json out.xlsx
入力以外のセルはすべて計算式。黄色のセルを変えると全シートが再計算される。
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

GOTHIC = "游ゴシック"
# 配色 B（生命保険）
MAIN, GOLD, LIGHT_SERIES, GRAY = "24513F", "B5852E", "A3BDAE", "9AA3AE"
TEXT, NOTE, RULE, INPUT, BAND = "333333", "4F5968", "C5CBD3", "FFF8E1", "EEF3EF"
YEN = '#,##0;[Red]-#,##0'
MAN = '#,##0;[Red]-#,##0'
PCT = '0.0%;[Red]-0.0%'
RATE = '0.00%'
DEC4 = '0.0000'
DEC2 = '0.00'

thin = Side(style="thin", color=RULE)


def F(**k):
    return Font(name=GOTHIC, size=k.pop("size", 10), color=k.pop("color", TEXT), **k)


AGE_MIN, AGE_MAX = 20, 109          # 推移のシートに並べる年齢
TBL_FIRST = 5                       # 基数シートの最初のデータ行（0歳）
TBL_LAST = TBL_FIRST + 110          # 110歳（生存数0）の行


def hdr(ws, row, values, col=1, height=42):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=col + i, value=v)
        c.font = F(bold=True, color="FFFFFF", size=9)
        c.fill = PatternFill("solid", fgColor=MAIN)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = height


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = F(size=14, bold=True, color=MAIN)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = F(size=9, color=NOTE)


def widths(ws, w):
    for i, x in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(i)].width = x


# ---------------------------------------------------------------- 前提
INPUTS = [
    # key, label, value, fmt, note
    ("S0", "基本保険金額（契約時）", 100_000_000, YEN, "円。設例"),
    ("X0", "加入年齢（男性）", 40, "0", "歳。20〜60歳の範囲で変更できる"),
    ("N", "保険料払込期間", 25, "0", "年。加入年齢＋払込期間＝払込満了年齢"),
    ("I", "予定利率", 0.0275, RATE, "2017年5月1日以後の平準払契約（特別勘定のデータ）。契約時期で異なる"),
    ("RP", "運用利回り（払込期間中）", 0.06, PCT, "特別勘定の利回り。運用関係費用を引いた後（仮定）"),
    ("RA", "運用利回り（払込満了後）", 0.06, PCT, "同上（仮定）"),
    ("FEE", "払込満了後の保険関係費用（積立金に対する年率）", 0.0, PCT, "会社は数値を開示していない。0%と仮定"),
    ("KS", "減額を始める年齢", 65, "0", "歳。払込満了年齢以後に限る"),
    ("KE", "減額を終える年齢", 99, "0", "歳。この年齢の契約応当日まで毎年減額する"),
    ("W", "一定額：毎年受け取る解約返戻金", 10_000_000, YEN, "円。積立金がこの額に満たない年は全部を解約"),
    ("K", "一定割合：毎年の減額割合", 0.10, PCT, "その時点の基本保険金額に対する割合"),
    ("MINS", "減額後の基本保険金額の最低額", 0, YEN, "円。会社所定（不明）。下回るときは全部を解約として計算"),
]
REF = {}


def build_inputs(wb):
    ws = wb.active
    ws.title = "前提"
    title(ws, "変額終身保険（プルデンシャル生命・終身型）　払込満了後の減額による推移の試算",
          "設例・年単位の簡略なモデル。黄色のセルを変えると、すべてのシートが計算し直される。金額の単位は円（比較のシートは万円）。")
    hdr(ws, 4, ["項目", "値", "説明"], height=20)
    r = 5
    for key, label, val, fmt, note in INPUTS:
        ws.cell(row=r, column=1, value=label).font = F()
        c = ws.cell(row=r, column=2, value=val)
        c.number_format = fmt
        c.fill = PatternFill("solid", fgColor=INPUT)
        c.font = F(bold=True, color=MAIN)
        c.alignment = Alignment(horizontal="right")
        ws.cell(row=r, column=3, value=note).font = F(size=9, color=NOTE)
        for col in range(1, 4):
            ws.cell(row=r, column=col).border = Border(bottom=thin)
        REF[key] = f"前提!$B${r}"
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="計算した値").font = F(bold=True, color=MAIN)
    r += 1
    x0, n, s0 = REF["X0"], REF["N"], REF["S0"]
    idx = lambda col, age: f"INDEX(基数!${col}${TBL_FIRST}:${col}${TBL_LAST},{age}+1)"
    derived = [
        ("XM", "払込満了年齢", f"={x0}+{n}", "0", "歳"),
        ("AX0", "一時払保険料率 A（加入年齢）", f"={idx('I', x0)}", DEC4, "基本保険金額1円当たりの終身保険の一時払純保険料"),
        ("ADUE", "期始払年金現価 ä（加入年齢・払込期間）",
         f"=({idx('H', x0)}-{idx('H', f'{x0}+{n}')})/{idx('E', x0)}", DEC4, "(N[x]−N[x+n])÷D[x]"),
        ("P", "年払純保険料", f"={s0}*B{r + 1}/B{r + 2}", YEN,
         "基本保険金額×A÷ä。付加保険料（会社の費用）を除く。積立金に入るのはこの額（仮定）"),
        ("PU", "年払純保険料（基本保険金額1円当たり）", f"=B{r + 3}/{s0}", "0.000000", ""),
        ("PT", "純保険料の払込総額", f"=B{r + 3}*{n}", YEN, ""),
    ]
    for key, label, formula, fmt, note in derived:
        ws.cell(row=r, column=1, value=label).font = F()
        c = ws.cell(row=r, column=2, value=formula)
        c.number_format = fmt
        c.alignment = Alignment(horizontal="right")
        ws.cell(row=r, column=3, value=note).font = F(size=9, color=NOTE)
        for col in range(1, 4):
            ws.cell(row=r, column=col).border = Border(bottom=thin)
        REF[key] = f"前提!$B${r}"
        r += 1

    r += 1
    notes = [
        ("計算の仕組み（約款・資料による）", True),
        ("・死亡保険金額＝基本保険金額＋変動保険金額。変動保険金額が負のときは基本保険金額（約款第6条1項）。", False),
        ("・変動保険金額は月単位の契約応当日ごとに計算する。積立金から「基本保険金額を支払うために必要な金額」を差し引いた額をもとにし、負にもなる（第7条1〜3項）。", False),
        ("・計算式：（月末の積立金−予定利率で運用された場合の積立金）÷ 計算時における一時払保険料相当額 ＝ 翌月の変動保険金額（特別勘定のデータ「契約月別運用実績表」の読み方）。", False),
        ("・基本保険金額を減額すると、同じ割合で変動保険金額も減額される。変動保険金額だけの減額はできない。減額部分は解約したものとして扱う（第24条2〜4項）。", False),
        ("・解約控除は、契約日から10年未満かつ払込期間中の解約・減額にだけかかる。払込満了後の解約返戻金は積立金の額（ご契約のしおり「変額保険のご契約にかかる諸費用」）。", False),
        ("", False),
        ("このモデルの前提（簡略化）", True),
        ("・年単位で計算する（実際は変動保険金額を毎月、積立金を毎日計算する）。保険料は年払・期始払、減額は各年の契約応当日に行う。", False),
        ("・予定死亡率と、積立金から差し引く危険保険料の死亡率は、生保標準生命表2018（死亡保険用・男）とする（実際の会社の率は不明）。", False),
        ("・「必要な金額」と「一時払保険料相当額」は、予定利率と上の死亡率による純保険料式で計算する（予定事業費を含めない）。", False),
        ("・積立金の推移：期末積立金 ＝ ｛期始積立金 ×（1＋運用利回り）− 死亡率 × 死亡保険金額｝÷（1 − 死亡率）。危険保険料は、死亡保険金額と積立金の差（危険保険金）に対してかかる。", False),
        ("・減額の受取額 ＝ 減額前の積立金 × 減額割合。一定額のときの減額割合 ＝ 受け取る額 ÷ 減額前の積立金。", False),
        ("・契約日から10年未満の積立金は解約控除を引く前の額（払込満了後は解約返戻金と同じ）。", False),
        ("", False),
        ("出典", True),
        ("・プルデンシャル生命「変額保険（終身型）普通保険約款」第3条・第6条・第7条・第22条・第24条、ご契約のしおり（NotebookLM「変額保険 ― 各社の約款・特別勘定のしおり」のソース、確認日 2026-10-15）", False),
        ("・プルデンシャル生命「ディスクロージャー資料 データ編　変額保険（終身型・有期型）（特別勘定）（2026年8月末）」9頁「契約月別運用実績表」の読み方、予定利率（同上）", False),
        ("・日本アクチュアリー会「生保標準生命表2018（死亡保険用）（男）」 https://www.actuaries.jp/lib/standard-life-table/pdf/seimeihyo2018.pdf （確認日 2026-10-15）", False),
    ]
    for text, bold in notes:
        c = ws.cell(row=r, column=1, value=text)
        c.font = F(bold=bold, color=MAIN if bold else NOTE, size=10 if bold else 9)
        r += 1
    widths(ws, [44, 18, 90])
    ws.freeze_panes = "A5"
    return ws


# ---------------------------------------------------------------- 基数
def build_table(wb, qx):
    ws = wb.create_sheet("基数")
    title(ws, "生命表と計算基数（予定利率による）",
          "q：生保標準生命表2018（死亡保険用・男）。l0＝100,000。D＝l×v^x、C＝d×v^(x+1)、M・N はその年齢から最終年齢までの合計、A＝M÷D、ä＝N÷D。")
    hdr(ws, 4, ["年齢 x", "死亡率 q", "生存数 l", "死亡数 d", "D", "C", "M", "N", "A（一時払保険料率）", "ä"], height=30)
    i = REF["I"]
    for age in range(0, 111):
        r = TBL_FIRST + age
        ws.cell(row=r, column=1, value=age)
        if age <= 109:
            c = ws.cell(row=r, column=2, value=qx[age])
            c.fill = PatternFill("solid", fgColor=INPUT)
        else:
            ws.cell(row=r, column=2, value=0)
        ws.cell(row=r, column=3, value=100000 if age == 0 else f"=C{r - 1}*(1-B{r - 1})")
        ws.cell(row=r, column=4, value=f"=C{r}*B{r}")
        ws.cell(row=r, column=5, value=f"=C{r}*(1+{i})^(-A{r})")
        ws.cell(row=r, column=6, value=f"=D{r}*(1+{i})^(-(A{r}+1))")
        ws.cell(row=r, column=7, value=f"=SUM(F{r}:F${TBL_LAST})")
        ws.cell(row=r, column=8, value=f"=SUM(E{r}:E${TBL_LAST})")
        ws.cell(row=r, column=9, value=f"=IF(E{r}>0,G{r}/E{r},0)")
        ws.cell(row=r, column=10, value=f"=IF(E{r}>0,H{r}/E{r},0)")
        fmts = ["0", "0.00000", "#,##0.0000", "#,##0.0000", "#,##0.0000", "#,##0.0000",
                "#,##0.0000", "#,##0.0000", DEC4, DEC4]
        for col, fm in enumerate(fmts, 1):
            cc = ws.cell(row=r, column=col)
            cc.number_format = fm
            cc.font = F(size=9)
    widths(ws, [8, 10, 12, 10, 12, 12, 12, 14, 12, 10])
    ws.freeze_panes = "B5"


# ---------------------------------------------------------------- 推移
COLS = ["年齢", "経過年数", "状態", "純保険料", "純保険料の累計", "一時払保険料率 A",
        "必要な金額（基本保険金額1円当たり）", "積立金（期始・減額前）", "基本保険金額（減額前）",
        "変動保険金額（減額前）", "減額割合", "受取額（減額部分の解約返戻金）", "受取額の累計",
        "基本保険金額（減額後）", "変動保険金額（減額後）", "積立金（減額後）＝解約返戻金",
        "死亡保険金額", "運用利回り", "死亡率 q", "期末積立金（生存者1件当たり）",
        "死亡保険金額÷解約返戻金", "変動保険金額÷基本保険金額", "受取累計＋解約返戻金"]
FIRST = 5   # 推移シートの最初のデータ行


def build_projection(wb, name, mode, subtitle):
    ws = wb.create_sheet(name)
    title(ws, name, subtitle)
    hdr(ws, 4, COLS)
    R = REF
    last = FIRST + (AGE_MAX - AGE_MIN)
    for k, age in enumerate(range(AGE_MIN, AGE_MAX + 1)):
        r = FIRST + k
        p = r - 1
        first = (k == 0)
        tbl = lambda col, a: f"INDEX(基数!${col}${TBL_FIRST}:${col}${TBL_LAST},{a}+1)"
        f = {}
        f["A"] = age
        f["B"] = f"=A{r}-{R['X0']}"
        f["C"] = f'=IF(B{r}<0,"加入前",IF(B{r}<{R["N"]},"払込中","払込満了後"))'
        f["D"] = f"=IF(AND(B{r}>=0,B{r}<{R['N']}),{R['P']},0)"
        f["E"] = f"=D{r}" if first else f"=E{p}+D{r}"
        f["F"] = f"={tbl('I', f'A{r}')}"
        # 必要な金額（1円当たり）：払込中は A − 純保険料率×(ä[x:n−t] − 1)、払込満了後は A
        m = f"({R['N']}-B{r})"
        adue = f"({tbl('H', f'A{r}')}-{tbl('H', f'A{r}+{m}')})/{tbl('E', f'A{r}')}"
        f["G"] = f"=IF(B{r}<0,0,IF(B{r}<{R['N']},F{r}-{R['PU']}*({adue}-1),F{r}))"
        f["H"] = f"=IF(B{r}<0,0,D{r})" if first else f"=IF(B{r}<0,0,T{p}+D{r})"
        f["I"] = f"=IF(B{r}<0,0,{R['S0']})" if first else f"=IF(B{r}<0,0,IF(B{r}=0,{R['S0']},N{p}))"
        f["J"] = f"=IF(AND(B{r}>=0,F{r}>0,I{r}>0),(H{r}-I{r}*G{r})/F{r},0)"
        cond = f"AND(B{r}>={R['N']},A{r}>={R['KS']},A{r}<={R['KE']},H{r}>0,I{r}>0)"
        if mode == "none":
            f["K"] = "=0"
        else:
            k0 = f"MIN(1,{R['W']}/H{r})" if mode == "amount" else f"{R['K']}"
            f["K"] = f"=IF({cond},IF(I{r}*(1-{k0})<{R['MINS']},1,{k0}),0)"
        f["L"] = f"=H{r}*K{r}"
        f["M"] = f"=L{r}" if first else f"=M{p}+L{r}"
        f["N"] = f"=I{r}*(1-K{r})"
        f["O"] = f"=J{r}*(1-K{r})"
        f["P"] = f"=H{r}*(1-K{r})"
        f["Q"] = f"=IF(N{r}>0,N{r}+MAX(0,O{r}),0)"
        f["R"] = f"=IF(B{r}<{R['N']},{R['RP']},{R['RA']}-{R['FEE']})"
        f["S"] = f"={tbl('B', f'A{r}')}"
        f["T"] = f"=IF(OR(B{r}<0,S{r}>=1,N{r}<=0),0,(P{r}*(1+R{r})-S{r}*Q{r})/(1-S{r}))"
        f["U"] = f'=IF(P{r}>0,Q{r}/P{r},0)'
        f["V"] = f'=IF(N{r}>0,O{r}/N{r},0)'
        f["W"] = f"=M{r}+P{r}"
        fmts = {"A": "0", "B": "0", "D": YEN, "E": YEN, "F": DEC4, "G": DEC4, "H": YEN, "I": YEN,
                "J": YEN, "K": PCT, "L": YEN, "M": YEN, "N": YEN, "O": YEN, "P": YEN, "Q": YEN,
                "R": PCT, "S": "0.00000", "T": YEN, "U": DEC2, "V": DEC2, "W": YEN}
        for col, val in f.items():
            c = ws[f"{col}{r}"]
            c.value = val
            c.font = F(size=9, bold=col in ("N", "O", "P"), color=MAIN if col in ("N", "O", "P") else TEXT)
            if col in fmts:
                c.number_format = fmts[col]
            if col in ("N", "O", "P", "Q"):
                c.fill = PatternFill("solid", fgColor=BAND)
            c.border = Border(bottom=thin)
    widths(ws, [6, 7, 10, 12, 14, 9, 11, 15, 15, 15, 8, 14, 15, 15, 15, 15, 15, 8, 9, 15, 9, 9, 15])
    ws.freeze_panes = "D5"

    return ws


# ---------------------------------------------------------------- 比較
SCEN = [("減額なし", "減額なし"), ("一定額で減額", "一定額"), ("一定割合で減額", "一定割合")]
ITEMS = [("N", "基本保険金額"), ("O", "変動保険金額"), ("Q", "死亡保険金額"), ("P", "解約返戻金"), ("M", "受取累計")]


def build_compare(wb):
    ws = wb.create_sheet("比較", 1)
    title(ws, "比較：減額なし・一定額・一定割合（単位：万円）",
          "各年齢の契約応当日の減額後の値。変動保険金額は負のこともある（死亡保険金額はそのとき基本保険金額）。グラフは減額する2つの方法だけを描く（減額なしは「減額なし」シートのグラフ）。")
    ws.cell(row=4, column=1, value="年齢")
    col = 2
    for sheet, short in SCEN:
        ws.cell(row=3, column=col, value=short)
        ws.merge_cells(start_row=3, start_column=col, end_row=3, end_column=col + len(ITEMS) - 1)
        c = ws.cell(row=3, column=col)
        c.font = F(bold=True, color=MAIN)
        c.alignment = Alignment(horizontal="center")
        col += len(ITEMS)
    hdr(ws, 4, ["年齢"] + [lab for _ in SCEN for _, lab in ITEMS], height=30)
    ages = [40, 50, 60, 65, 66, 67, 70, 75, 77, 80, 85, 90, 95, 100]
    for k, age in enumerate(ages):
        r = 5 + k
        ws.cell(row=r, column=1, value=age).font = F(bold=True)
        col = 2
        for sheet, _ in SCEN:
            for item, _ in ITEMS:
                fml = (f"=INDEX('{sheet}'!${item}${FIRST}:${item}${FIRST + AGE_MAX - AGE_MIN},"
                       f"MATCH($A{r},'{sheet}'!$A${FIRST}:$A${FIRST + AGE_MAX - AGE_MIN},0))/10000")
                c = ws.cell(row=r, column=col, value=fml)
                c.number_format = MAN
                c.font = F(size=9)
                c.border = Border(bottom=thin)
                if sheet != "減額なし" and item == "P":
                    c.fill = PatternFill("solid", fgColor=BAND)
                col += 1
        ws.cell(row=r, column=1).border = Border(bottom=thin)
    widths(ws, [7] + [11] * (len(SCEN) * len(ITEMS)))
    ws.freeze_panes = "B5"



# ---------------------------------------------------------------- グラフ用（万円）
CH_A0, CH_A1 = 55, 100              # グラフに描く年齢
CH_FIRST = 5
CH_ITEMS = [("N", "基本保険金額"), ("O", "変動保険金額"), ("Q", "死亡保険金額"),
            ("P", "解約返戻金"), ("M", "受取累計")]
CH_COL = {}                         # (シート名, 項目) → グラフ用シートの列番号


def build_chartdata(wb):
    ws = wb.create_sheet("グラフ用")
    title(ws, "グラフ用のデータ（単位：万円）",
          "推移の各シートから、各年齢の契約応当日の減額後の値を万円に直して並べる。グラフはこのシートを参照する。")
    ws.cell(row=4, column=1, value="年齢")
    col = 2
    for sheet, short in SCEN:
        ws.cell(row=3, column=col, value=short).font = F(bold=True, color=MAIN)
        for item, _ in CH_ITEMS:
            CH_COL[(sheet, item)] = col
            col += 1
    hdr(ws, 4, ["年齢"] + [lab for _ in SCEN for _, lab in CH_ITEMS], height=30)
    last = FIRST + AGE_MAX - AGE_MIN
    for k, age in enumerate(range(CH_A0, CH_A1 + 1)):
        r = CH_FIRST + k
        ws.cell(row=r, column=1, value=age).font = F(size=9)
        for (sheet, item), c in CH_COL.items():
            fml = (f"=INDEX('{sheet}'!${item}${FIRST}:${item}${last},"
                   f"MATCH($A{r},'{sheet}'!$A${FIRST}:$A${last},0))/10000")
            cell = ws.cell(row=r, column=c, value=fml)
            cell.number_format = MAN
            cell.font = F(size=9)
    widths(ws, [7] + [11] * (len(SCEN) * len(CH_ITEMS)))
    ws.freeze_panes = "B5"


def line_chart(title_text, height=8, width=17):
    ch = LineChart()
    ch.title = title_text
    ch.height, ch.width = height, width
    ch.y_axis.title = "万円"
    ch.y_axis.number_format = '#,##0'
    ch.x_axis.title = "年齢（各年の契約応当日の減額後）"
    ch.x_axis.tickLblSkip = 5
    # openpyxl 3.1 は軸の delete を書き出さず、Excel で目盛りが消えることがあるため明示する
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.legend.position = "b"
    return ch


def add_series(ch, ws, col, label, color, width_pt, dash=None):
    from openpyxl.chart.series import SeriesLabel
    n = CH_A1 - CH_A0
    ch.add_data(Reference(ws, min_col=col, min_row=CH_FIRST, max_row=CH_FIRST + n), titles_from_data=False)
    s = ch.series[-1]
    s.tx = SeriesLabel(v=label)
    s.graphicalProperties.line.solidFill = color
    s.graphicalProperties.line.width = int(width_pt * 12700)
    if dash:
        s.graphicalProperties.line.dashStyle = dash
    s.smooth = False


def add_charts(wb):
    cd = wb["グラフ用"]
    n = CH_A1 - CH_A0
    cats = Reference(cd, min_col=1, min_row=CH_FIRST, max_row=CH_FIRST + n)
    # 推移の各シート：基本保険金額・変動保険金額・解約返戻金・死亡保険金額
    for sheet, _ in SCEN:
        ch = line_chart(f"{sheet}：保険金額と解約返戻金", height=9, width=20)
        for item, label, color, w, dash in [("Q", "死亡保険金額", GRAY, 1.5, "dash"),
                                            ("N", "基本保険金額", LIGHT_SERIES, 2.25, None),
                                            ("O", "変動保険金額", GOLD, 2.25, None),
                                            ("P", "解約返戻金（積立金）", MAIN, 3, None)]:
            add_series(ch, cd, CH_COL[(sheet, item)], label, color, w, dash)
        ch.set_categories(cats)
        wb[sheet].add_chart(ch, "Y5")
    # 比較シート：3つのグラフを縦に並べる（横に並べると幅が列幅を超えて重なる）
    ws = wb["比較"]
    colors = {"一定額で減額": MAIN, "一定割合で減額": GOLD}
    anchor_row = 5 + 14 + 2
    for j, (item, label) in enumerate([("P", "解約返戻金（積立金）"), ("Q", "死亡保険金額"), ("M", "受取額の累計")]):
        # 減額なし（100歳で5億円超）を入れると目盛りが広がり、減額する2つの方法の差が読めなくなるため除く
        ch = line_chart(f"{label}：一定額と一定割合")
        for sheet, short in SCEN:
            if sheet == "減額なし":
                continue
            add_series(ch, cd, CH_COL[(sheet, item)], short, colors[sheet], 3)
        ch.set_categories(cats)
        ws.add_chart(ch, f"A{anchor_row + j * 18}")


def main(qx_path, out):
    qx = {int(k): v for k, v in json.load(open(qx_path)).items()}
    wb = Workbook()
    build_inputs(wb)
    build_table(wb, qx)
    build_projection(wb, "減額なし", "none", "払込満了後も減額せずに保有する場合（比較の基準）。")
    build_projection(wb, "一定額で減額", "amount",
                     "減額を始める年齢から毎年、決めた額の解約返戻金を受け取れるだけ基本保険金額を減額する。")
    build_projection(wb, "一定割合で減額", "ratio",
                     "減額を始める年齢から毎年、その時点の基本保険金額を決めた割合だけ減額する。")
    build_compare(wb)
    build_chartdata(wb)
    add_charts(wb)
    wb.save(out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
